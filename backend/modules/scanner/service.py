from pathlib import Path

from modules.scanner.malicious_url import scan_urls
from modules.scanner.prompt_injection import scan_text
from modules.scanner.sensitive_data import scan_sensitive_data


def read_text_file(path: str) -> str:
    suffix = Path(path).suffix.lower()
    if suffix in (".txt", ".md", ".csv", ".json", ".log"):
        return Path(path).read_text(encoding="utf-8", errors="ignore")
    if suffix == ".pdf":
        return _read_pdf(path)
    if suffix == ".docx":
        return _read_docx(path)
    return Path(path).read_bytes()[:20000].decode("utf-8", errors="ignore")


def _read_pdf(path: str) -> str:
    try:
        from pypdf import PdfReader

        reader = PdfReader(path)
        chunks = []
        for idx, page in enumerate(reader.pages[:30], start=1):
            text = page.extract_text() or ""
            if text.strip():
                chunks.append(f"[page {idx}]\n{text}")
        return "\n\n".join(chunks)
    except Exception as exc:
        return f"[pdf_extract_error] {exc}\n" + Path(path).read_bytes()[:20000].decode("utf-8", errors="ignore")


def _read_docx(path: str) -> str:
    try:
        from docx import Document

        doc = Document(path)
        chunks = []
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                chunks.append(paragraph.text)
        for table_idx, table in enumerate(doc.tables, start=1):
            chunks.append(f"[table {table_idx}]")
            for row in table.rows:
                chunks.append(" | ".join(cell.text.strip() for cell in row.cells))
        return "\n".join(chunks)
    except Exception as exc:
        return f"[docx_extract_error] {exc}\n" + Path(path).read_bytes()[:20000].decode("utf-8", errors="ignore")


def scan_file(path: str) -> dict:
    text = read_text_file(path)
    prompt_result = scan_text(text)
    findings = list(prompt_result["findings"])
    findings.extend(scan_sensitive_data(text))
    findings.extend(scan_urls(text))

    score = prompt_result["risk_score"]
    if any(f["category"] == "sensitive_leak" for f in findings):
        score = max(score, 50)
    if any(f["category"] == "malicious_url" for f in findings):
        score = max(score, 45)

    from modules.policy.engine import risk_level

    return {"risk_score": score, "risk_level": risk_level(score), "findings": findings}
